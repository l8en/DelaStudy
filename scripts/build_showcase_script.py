"""Build the timed companion DOCX with the managed Codex Python runtime.

Design: compact_reference_guide; workshop_agenda title stack, without a table.
Named overrides: 24pt title, 10pt subtitle, 9pt screen cues, 9pt page furniture.
All are explicit and reused. No narration is composited into the video.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def main():
    source, target = map(Path, sys.argv[1:3])
    data = json.loads(source.read_text())
    scenes = data["scenes"]
    assert scenes[0]["start"] == 0 and scenes[-1]["end"] == 135
    assert all(a["end"] == b["start"] for a, b in zip(scenes, scenes[1:]))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    def style(name, size, color, before, after, bold=False):
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = False
        s.font.color.rgb = RGBColor.from_string(color)
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            s.element.get_or_add_rPr().get_or_add_rFonts().set(qn(f"w:{script}"), "Calibri")
        p = s.paragraph_format
        p.space_before, p.space_after = Pt(before), Pt(after)
        p.line_spacing = 1.25
        p.widow_control = True
        p.keep_with_next = name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Screen cue")
        for border in list(s.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        return s

    style("Normal", 11, "172432", 0, 6)
    style("Title", 24, "172432", 0, 4, True)
    style("Subtitle", 10, "425469", 0, 10)
    style("Heading 1", 16, "2E74B5", 18, 10, True)
    style("Heading 2", 13, "2E74B5", 14, 7, True)
    style("Heading 3", 12, "1F4D78", 10, 5, True)
    style("Screen cue", 9, "425469", 0, 5)
    style("Header", 9, "425469", 0, 0)
    style("Footer", 9, "425469", 0, 0)
    section.header.paragraphs[0].text = "DelaStudy  |  Timed narration script"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("2:15 walkthrough  •  ")
    page = OxmlElement("w:fldSimple"); page.set(qn("w:instr"), "PAGE"); footer._p.append(page)

    def stamp(seconds):
        return f"{seconds // 60}:{seconds % 60:02}"

    page_starts = {0: ("DelaStudy", "M’hamed  |  0:00–0:41  |  Opening & setup"),
                   5: ("Learn, then use it", "Presenter 2  |  0:41–1:36  |  Lessons & planning"),
                   10: ("Focus and follow through", "Presenter 3  |  1:36–2:15  |  Sound, notification & progress")}
    if "pageStarts" in data:
        page_starts = {item["index"]: (item["title"], item["subtitle"]) for item in data["pageStarts"]}
    for i, scene in enumerate(scenes):
        if i in page_starts:
            if i: doc.add_page_break()
            title, subtitle = page_starts[i]
            doc.add_paragraph(title, "Title")
            doc.add_paragraph(subtitle, "Subtitle")
            if i == 0:
                doc.add_paragraph(data.get("instruction", "Read the body text aloud. The smaller screen cues are not narration. Start each block at its timestamp; leave a brief pause if you finish early."), "Screen cue")
        doc.add_paragraph(f"{stamp(scene['start'])}–{stamp(scene['end'])}  {scene['title']}", "Heading 2")
        doc.add_paragraph("ON SCREEN  " + scene["screen"], "Screen cue")
        if "segments" in scene:
            for segment in scene["segments"]:
                p = doc.add_paragraph(style="Normal")
                cue = p.add_run(f"{stamp(segment['start'])}–{stamp(segment['end'])}  ")
                cue.font.size = Pt(9); cue.font.color.rgb = RGBColor.from_string("425469")
                p.add_run(segment["read"])
        else:
            doc.add_paragraph(scene["read"], "Normal")
    doc.add_paragraph("Recording note: actual iPhone Simulator footage in an iPhone 14 Pro-style frame. One real local Focus notification; the one-minute wait is shortened in the edit. The video is silent. Accounts and progress are local to the device.", "Screen cue")
    doc.core_properties.title = "DelaStudy — 2:15 walkthrough narration"
    doc.core_properties.author = data.get("author", "M’hamed")
    doc.core_properties.subject = "Time-aligned companion script; not video captions"
    doc.save(target)
    words = sum(len(s["read"].split()) for s in scenes)
    print(json.dumps({"output": str(target), "spokenWords": words, "duration": 135, "preset": "compact_reference_guide"}))


if __name__ == "__main__":
    main()
